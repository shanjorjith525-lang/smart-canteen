from datetime import datetime, timedelta

from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import cm

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from .models import Food, Order, FoodWaste


# =========================================================
# HOME / FOOD LIST
# =========================================================

def food_list(request):
    foods = Food.objects.all().order_by('name')

    return render(
        request,
        'menu/food_list.html',
        {
            'foods': foods,
        }
    )


# =========================================================
# REGISTER
# =========================================================

def register(request):

    if request.user.is_authenticated:
        return redirect('food_list')

    if request.method == 'POST':

        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')
        password2 = request.POST.get('password2', '')

        if not username or not password or not password2:
            messages.error(request, 'Please fill all fields.')
            return render(request, 'menu/register.html')

        if password != password2:
            messages.error(request, 'Passwords do not match.')
            return render(request, 'menu/register.html')

        if User.objects.filter(username=username).exists():
            messages.error(request, 'Username already exists.')
            return render(request, 'menu/register.html')

        user = User.objects.create_user(
            username=username,
            password=password
        )

        login(request, user)

        messages.success(request, 'Registration successful!')

        return redirect('food_list')

    return render(request, 'menu/register.html')


# =========================================================
# LOGIN
# =========================================================

def user_login(request):

    if request.user.is_authenticated:
        return redirect('food_list')

    if request.method == 'POST':

        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None:

            login(request, user)

            # Staff users can still access the normal home page.
            return redirect('food_list')

        messages.error(
            request,
            'Invalid username or password.'
        )

    return render(request, 'menu/login.html')


# =========================================================
# LOGOUT
# =========================================================

@login_required
def user_logout(request):

    logout(request)

    return redirect('food_list')


# =========================================================
# CUSTOMER DASHBOARD
# =========================================================

@login_required
def customer_dashboard(request):

    orders = Order.objects.filter(
        customer=request.user
    ).select_related('food')

    total_orders = orders.count()

    pending_orders = orders.filter(
        status='pending'
    ).count()

    preparing_orders = orders.filter(
        status='preparing'
    ).count()

    ready_orders = orders.filter(
        status='ready'
    ).count()

    completed_orders = orders.filter(
        status='completed'
    ).count()

    total_spent = sum(
        order.total_price
        for order in orders
    )

    recent_orders = orders.order_by(
        '-ordered_at'
    )[:5]

    return render(
        request,
        'menu/customer_dashboard.html',
        {
            'total_orders': total_orders,
            'pending_orders': pending_orders,
            'preparing_orders': preparing_orders,
            'ready_orders': ready_orders,
            'completed_orders': completed_orders,
            'total_spent': total_spent,
            'recent_orders': recent_orders,
        }
    )


# =========================================================
# ORDER FOOD
# =========================================================

@login_required
def order_food(request, food_id):

    food = get_object_or_404(
        Food,
        id=food_id
    )

    if request.method == 'POST':

        try:
            quantity = int(
                request.POST.get('quantity', '1')
            )
        except (ValueError, TypeError):
            quantity = 0

        if quantity <= 0:

            messages.error(
                request,
                'Please enter a valid quantity.'
            )

            return render(
                request,
                'menu/order.html',
                {'food': food}
            )

        if quantity > food.available_quantity:

            messages.error(
                request,
                f'Only {food.available_quantity} item(s) available.'
            )

            return render(
                request,
                'menu/order.html',
                {'food': food}
            )

        total_price = food.price * quantity

        order = Order.objects.create(
            customer=request.user,
            food=food,
            quantity=quantity,
            total_price=total_price,
            status='pending'
        )

        food.available_quantity -= quantity

        if food.available_quantity == 0:
            food.is_available = False

        food.save()

        return render(
            request,
            'menu/order_success.html',
            {
                'order': order,
            }
        )

    return render(
        request,
        'menu/order.html',
        {
            'food': food,
        }
    )


# =========================================================
# MY ORDERS
# =========================================================

@login_required
def my_orders(request):

    orders = (
        Order.objects
        .filter(customer=request.user)
        .select_related('food')
        .order_by('-ordered_at')
    )

    return render(
        request,
        'menu/my_orders.html',
        {
            'orders': orders,
        }
    )


# =========================================================
# ORDER STATUS
# =========================================================

@login_required
def order_status(request, order_id):

    order = get_object_or_404(
        Order.objects.select_related('food'),
        id=order_id,
        customer=request.user
    )

    return render(
        request,
        'menu/order_status.html',
        {
            'order': order,
        }
    )


# =========================================================
# ANALYTICS DASHBOARD
# =========================================================

@login_required
def analytics_dashboard(request):

    if not request.user.is_staff:
        return redirect('food_list')

    today = timezone.localdate()

    today_orders = Order.objects.filter(
        ordered_at__date=today
    )

    completed_orders = today_orders.filter(
        status='completed'
    )

    food_waste = FoodWaste.objects.filter(
        date=today
    )

    total_orders = today_orders.count()

    food_consumed = (
        completed_orders.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_waste = (
        food_waste.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_sales = (
        today_orders.aggregate(
            total=Sum('total_price')
        )['total'] or 0
    )

    food_consumption = (
        completed_orders
        .values('food__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
    )

    food_waste_data = (
        food_waste
        .values('food__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
    )

    recent_waste = (
        FoodWaste.objects
        .select_related('food')
        .order_by('-created_at')[:10]
    )

    return render(
        request,
        'menu/analytics_dashboard.html',
        {
            'today': today,
            'total_orders': total_orders,
            'food_consumed': food_consumed,
            'total_waste': total_waste,
            'total_sales': total_sales,
            'food_consumption': food_consumption,
            'food_waste': food_waste_data,
            'recent_waste': recent_waste,
        }
    )


# =========================================================
# ADMIN DASHBOARD
# =========================================================

@login_required
def admin_dashboard(request):

    if not request.user.is_staff:
        return redirect('food_list')

    today = timezone.localdate()

    orders = Order.objects.all()

    completed_orders = orders.filter(
        status='completed'
    )

    food_waste = FoodWaste.objects.all()

    total_orders = orders.count()

    food_consumed = (
        completed_orders.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_waste = (
        food_waste.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_sales = (
        orders.aggregate(
            total=Sum('total_price')
        )['total'] or 0
    )

    food_consumption = (
        completed_orders
        .values('food__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
    )

    food_waste_data = (
        food_waste
        .values('food__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
    )

    recent_orders = (
        Order.objects
        .select_related('food', 'customer')
        .order_by('-ordered_at')[:10]
    )

    recent_waste = (
        FoodWaste.objects
        .select_related('food')
        .order_by('-created_at')[:10]
    )

    low_stock_foods = []

    for food in Food.objects.all():

        completed_food_orders = Order.objects.filter(
            food=food,
            status='completed'
        )

        total_quantity = (
            completed_food_orders.aggregate(
                total=Sum('quantity')
            )['total'] or 0
        )

        total_days = (
            completed_food_orders
            .dates('ordered_at', 'day')
            .count()
        )

        if total_days > 0:
            average_daily_demand = (
                total_quantity / total_days
            )
        else:
            average_daily_demand = 0

        recommended_quantity = round(
            average_daily_demand * 1.20
        )

        if food.available_quantity < recommended_quantity:

            low_stock_foods.append({
                'food': food,
                'current_stock': food.available_quantity,
                'recommended_quantity': recommended_quantity,
            })

    return render(
        request,
        'menu/admin_dashboard.html',
        {
            'today': today,
            'total_orders': total_orders,
            'food_consumed': food_consumed,
            'total_waste': total_waste,
            'total_sales': total_sales,
            'food_consumption': food_consumption,
            'food_waste': food_waste_data,
            'recent_orders': recent_orders,
            'recent_waste': recent_waste,
            'low_stock_foods': low_stock_foods,
        }
    )


# =========================================================
# DATE REPORT
# =========================================================

@login_required
def date_report(request):

    if not request.user.is_staff:
        return redirect('food_list')

    selected_date = request.GET.get('date')

    if selected_date:

        try:
            report_date = datetime.strptime(
                selected_date,
                '%Y-%m-%d'
            ).date()
        except ValueError:
            report_date = timezone.localdate()

    else:
        report_date = timezone.localdate()

    orders = Order.objects.filter(
        ordered_at__date=report_date
    )

    completed_orders = orders.filter(
        status='completed'
    )

    waste_records = FoodWaste.objects.filter(
        date=report_date
    )

    total_orders = orders.count()

    food_consumed = (
        completed_orders.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_waste = (
        waste_records.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_sales = (
        orders.aggregate(
            total=Sum('total_price')
        )['total'] or 0
    )

    food_consumption = (
        completed_orders
        .values('food__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
    )

    food_waste = (
        waste_records
        .values('food__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
    )

    chart_data = {}

    for item in food_consumption:

        name = item['food__name']

        chart_data[name] = {
            'consumption': item['total'],
            'waste': 0,
        }

    for item in food_waste:

        name = item['food__name']

        if name not in chart_data:

            chart_data[name] = {
                'consumption': 0,
                'waste': item['total'],
            }

        else:

            chart_data[name]['waste'] = item['total']

    chart_labels = list(chart_data.keys())

    chart_consumption = [
        chart_data[name]['consumption']
        for name in chart_labels
    ]

    chart_waste = [
        chart_data[name]['waste']
        for name in chart_labels
    ]

    return render(
        request,
        'menu/date_report.html',
        {
            'report_date': report_date,
            'total_orders': total_orders,
            'food_consumed': food_consumed,
            'total_waste': total_waste,
            'total_sales': total_sales,
            'food_consumption': food_consumption,
            'food_waste': food_waste,
            'chart_labels': chart_labels,
            'chart_consumption': chart_consumption,
            'chart_waste': chart_waste,
        }
    )


# =========================================================
# PDF REPORT
# =========================================================

@login_required
def pdf_report(request):

    if not request.user.is_staff:
        return redirect('food_list')

    selected_date = request.GET.get('date')

    if selected_date:

        try:
            report_date = datetime.strptime(
                selected_date,
                '%Y-%m-%d'
            ).date()
        except ValueError:
            report_date = timezone.localdate()

    else:
        report_date = timezone.localdate()

    orders = (
        Order.objects
        .filter(ordered_at__date=report_date)
        .select_related('food', 'customer')
    )

    waste_records = (
        FoodWaste.objects
        .filter(date=report_date)
        .select_related('food')
    )

    completed_orders = orders.filter(
        status='completed'
    )

    total_orders = orders.count()

    food_consumed = (
        completed_orders.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_waste = (
        waste_records.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_sales = (
        orders.aggregate(
            total=Sum('total_price')
        )['total'] or 0
    )

    response = HttpResponse(
        content_type='application/pdf'
    )

    response['Content-Disposition'] = (
        f'attachment; '
        f'filename="smart_canteen_{report_date}.pdf"'
    )

    pdf = canvas.Canvas(
        response,
        pagesize=A4
    )

    width, height = A4

    y = height - 2 * cm

    pdf.setFont(
        'Helvetica-Bold',
        20
    )

    pdf.drawString(
        2 * cm,
        y,
        'Smart Canteen'
    )

    y -= 0.8 * cm

    pdf.setFont(
        'Helvetica',
        11
    )

    pdf.drawString(
        2 * cm,
        y,
        f'Daily Report: {report_date}'
    )

    y -= 1.2 * cm

    pdf.setFont(
        'Helvetica-Bold',
        11
    )

    pdf.drawString(
        2 * cm,
        y,
        f'Total Orders: {total_orders}'
    )

    y -= 0.6 * cm

    pdf.drawString(
        2 * cm,
        y,
        f'Food Consumed: {food_consumed}'
    )

    y -= 0.6 * cm

    pdf.drawString(
        2 * cm,
        y,
        f'Food Waste: {total_waste}'
    )

    y -= 0.6 * cm

    pdf.drawString(
        2 * cm,
        y,
        f'Total Sales: Rs. {total_sales}'
    )

    y -= 1 * cm

    pdf.setFont(
        'Helvetica-Bold',
        12
    )

    pdf.drawString(
        2 * cm,
        y,
        'Orders'
    )

    y -= 0.6 * cm

    pdf.setFont(
        'Helvetica',
        9
    )

    for order in orders:

        customer = (
            order.customer.username
            if order.customer
            else 'Guest'
        )

        text = (
            f'#{order.id} | '
            f'{order.food.name} | '
            f'Qty: {order.quantity} | '
            f'Rs. {order.total_price} | '
            f'{order.status} | '
            f'{customer}'
        )

        pdf.drawString(
            2 * cm,
            y,
            text[:110]
        )

        y -= 0.45 * cm

        if y < 2 * cm:

            pdf.showPage()

            y = height - 2 * cm

            pdf.setFont(
                'Helvetica',
                9
            )

    y -= 0.5 * cm

    pdf.setFont(
        'Helvetica-Bold',
        12
    )

    pdf.drawString(
        2 * cm,
        y,
        'Food Waste'
    )

    y -= 0.6 * cm

    pdf.setFont(
        'Helvetica',
        9
    )

    for waste in waste_records:

        text = (
            f'{waste.food.name} | '
            f'Qty: {waste.quantity} | '
            f'Reason: {waste.get_reason_display()}'
        )

        pdf.drawString(
            2 * cm,
            y,
            text[:110]
        )

        y -= 0.45 * cm

        if y < 2 * cm:

            pdf.showPage()

            y = height - 2 * cm

            pdf.setFont(
                'Helvetica',
                9
            )

    pdf.save()

    return response


# =========================================================
# EXCEL REPORT
# =========================================================

@login_required
def excel_report(request):

    if not request.user.is_staff:
        return redirect('food_list')

    selected_date = request.GET.get('date')

    if selected_date:

        try:
            report_date = datetime.strptime(
                selected_date,
                '%Y-%m-%d'
            ).date()
        except ValueError:
            report_date = timezone.localdate()

    else:
        report_date = timezone.localdate()

    orders = (
        Order.objects
        .filter(ordered_at__date=report_date)
        .select_related('food', 'customer')
    )

    waste_records = (
        FoodWaste.objects
        .filter(date=report_date)
        .select_related('food')
    )

    completed_orders = orders.filter(
        status='completed'
    )

    total_orders = orders.count()

    food_consumed = (
        completed_orders.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_waste = (
        waste_records.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_sales = (
        orders.aggregate(
            total=Sum('total_price')
        )['total'] or 0
    )

    workbook = Workbook()

    # Daily Report
    sheet = workbook.active
    sheet.title = 'Daily Report'

    sheet.append([
        'Smart Canteen Daily Report'
    ])

    sheet.append([
        'Report Date',
        str(report_date)
    ])

    sheet.append([
        'Total Orders',
        total_orders
    ])

    sheet.append([
        'Food Consumed',
        food_consumed
    ])

    sheet.append([
        'Food Waste',
        total_waste
    ])

    sheet.append([
        'Total Sales',
        float(total_sales)
    ])

    # Orders
    orders_sheet = workbook.create_sheet(
        'Orders'
    )

    orders_sheet.append([
        'Order ID',
        'Customer',
        'Food',
        'Quantity',
        'Total Price',
        'Status',
        'Ordered At',
    ])

    for order in orders:

        customer = (
            order.customer.username
            if order.customer
            else 'Guest'
        )

        orders_sheet.append([
            order.id,
            customer,
            order.food.name,
            order.quantity,
            float(order.total_price),
            order.status,
            order.ordered_at.strftime(
                '%Y-%m-%d %H:%M'
            ),
        ])

    # Food Waste
    waste_sheet = workbook.create_sheet(
        'Food Waste'
    )

    waste_sheet.append([
        'Food',
        'Quantity',
        'Reason',
        'Date',
        'Notes',
    ])

    for waste in waste_records:

        waste_sheet.append([
            waste.food.name,
            waste.quantity,
            waste.get_reason_display(),
            str(waste.date),
            waste.notes,
        ])

    # Food Analysis
    analysis_sheet = workbook.create_sheet(
        'Food Analysis'
    )

    analysis_sheet.append([
        'Food',
        'Consumed',
        'Wasted',
    ])

    food_names = set()

    for order in completed_orders:
        food_names.add(order.food.name)

    for waste in waste_records:
        food_names.add(waste.food.name)

    for food_name in sorted(food_names):

        consumed = (
            completed_orders
            .filter(food__name=food_name)
            .aggregate(total=Sum('quantity'))['total']
            or 0
        )

        wasted = (
            waste_records
            .filter(food__name=food_name)
            .aggregate(total=Sum('quantity'))['total']
            or 0
        )

        analysis_sheet.append([
            food_name,
            consumed,
            wasted,
        ])

    # Excel formatting
    for worksheet in workbook.worksheets:

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal='center'
            )

        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                try:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

                except Exception:
                    pass

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 3,
                40
            )

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument'
            '.spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        f'attachment; '
        f'filename="smart_canteen_{report_date}.xlsx"'
    )

    workbook.save(response)

    return response


# =========================================================
# WEEKLY / MONTHLY ANALYTICS
# =========================================================

@login_required
def period_analytics(request):

    if not request.user.is_staff:
        return redirect('food_list')

    period = request.GET.get(
        'period',
        'week'
    )

    today = timezone.localdate()

    if period == 'month':

        start_date = today.replace(day=1)
        period_name = 'Monthly Analytics'

    else:

        start_date = (
            today - timedelta(
                days=today.weekday()
            )
        )

        period = 'week'
        period_name = 'Weekly Analytics'

    orders = Order.objects.filter(
        ordered_at__date__gte=start_date,
        ordered_at__date__lte=today
    )

    completed_orders = orders.filter(
        status='completed'
    )

    waste_records = FoodWaste.objects.filter(
        date__gte=start_date,
        date__lte=today
    )

    total_orders = orders.count()

    food_consumed = (
        completed_orders.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_waste = (
        waste_records.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    total_sales = (
        orders.aggregate(
            total=Sum('total_price')
        )['total'] or 0
    )

    total_handled = food_consumed + total_waste

    if total_handled > 0:

        waste_percentage = (
            total_waste / total_handled
        ) * 100

    else:

        waste_percentage = 0

    most_ordered = (
        completed_orders
        .values('food__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
        .first()
    )

    most_wasted = (
        waste_records
        .values('food__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
        .first()
    )

    food_consumption = (
        completed_orders
        .values('food__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
    )

    food_waste = (
        waste_records
        .values('food__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
    )

    chart_data = {}

    for item in food_consumption:

        name = item['food__name']

        chart_data[name] = {
            'consumption': item['total'],
            'waste': 0,
        }

    for item in food_waste:

        name = item['food__name']

        if name not in chart_data:

            chart_data[name] = {
                'consumption': 0,
                'waste': item['total'],
            }

        else:

            chart_data[name]['waste'] = item['total']

    chart_labels = list(
        chart_data.keys()
    )

    chart_consumption = [
        chart_data[name]['consumption']
        for name in chart_labels
    ]

    chart_waste = [
        chart_data[name]['waste']
        for name in chart_labels
    ]

    return render(
        request,
        'menu/period_analytics.html',
        {
            'today': today,
            'start_date': start_date,
            'period': period,
            'period_name': period_name,
            'total_orders': total_orders,
            'food_consumed': food_consumed,
            'total_waste': total_waste,
            'total_sales': total_sales,
            'waste_percentage': round(
                waste_percentage,
                2
            ),
            'most_ordered': most_ordered,
            'most_wasted': most_wasted,
            'food_consumption': food_consumption,
            'food_waste': food_waste,
            'chart_labels': chart_labels,
            'chart_consumption': chart_consumption,
            'chart_waste': chart_waste,
        }
    )


# =========================================================
# DEMAND PREDICTION
# =========================================================

@login_required
def demand_prediction(request):

    if not request.user.is_staff:
        return redirect('food_list')

    predictions = []

    for food in Food.objects.all():

        completed_orders = Order.objects.filter(
            food=food,
            status='completed'
        )

        total_quantity = (
            completed_orders.aggregate(
                total=Sum('quantity')
            )['total'] or 0
        )

        total_days = (
            completed_orders
            .dates('ordered_at', 'day')
            .count()
        )

        if total_days > 0:

            average_daily_demand = (
                total_quantity / total_days
            )

        else:

            average_daily_demand = 0

        recommended_quantity = round(
            average_daily_demand * 1.20
        )

        current_stock = food.available_quantity

        if current_stock < recommended_quantity:
            stock_status = 'Low Stock'
        else:
            stock_status = 'Stock Available'

        predictions.append({
            'food': food,
            'total_quantity': total_quantity,
            'total_days': total_days,
            'average_daily_demand': round(
                average_daily_demand,
                2
            ),
            'recommended_quantity':
                recommended_quantity,
            'current_stock':
                current_stock,
            'stock_status':
                stock_status,
        })

    return render(
        request,
        'menu/demand_prediction.html',
        {
            'predictions': predictions,
        }
    )


# =========================================================
# WASTE ANALYSIS
# =========================================================

@login_required
def waste_analysis(request):

    if not request.user.is_staff:
        return redirect('food_list')

    today = timezone.localdate()

    waste_records = FoodWaste.objects.all()

    total_waste = (
        waste_records.aggregate(
            total=Sum('quantity')
        )['total'] or 0
    )

    waste_by_reason = (
        waste_records
        .values('reason')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
    )

    reason_dict = dict(
        FoodWaste.REASON_CHOICES
    )

    chart_labels = []

    chart_values = []

    for item in waste_by_reason:

        chart_labels.append(
            reason_dict.get(
                item['reason'],
                item['reason']
            )
        )

        chart_values.append(
            item['total']
        )

    food_waste = (
        waste_records
        .values('food__name')
        .annotate(total=Sum('quantity'))
        .order_by('-total')
    )

    most_wasted = food_waste.first()

    most_common_reason = waste_by_reason.first()

    if most_common_reason:

        most_common_reason_name = reason_dict.get(
            most_common_reason['reason'],
            most_common_reason['reason']
        )

    else:

        most_common_reason_name = 'No Data'

    return render(
        request,
        'menu/waste_analysis.html',
        {
            'today': today,
            'total_waste': total_waste,
            'waste_by_reason': waste_by_reason,
            'food_waste': food_waste,
            'most_wasted': most_wasted,
            'most_common_reason':
                most_common_reason_name,
            'chart_labels': chart_labels,
            'chart_values': chart_values,
        }
    )